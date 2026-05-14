#!/usr/bin/env python3
"""
Erstellt aus den Markdown-Kapiteln ein druckfertiges DOCX.

Verwendung:
    python3 build_book_docx.py

Ausgabe:
    04_Export/Fremdfirmenkoordination.docx
"""

import os
import re
from pathlib import Path
import pypandoc
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

REPO_ROOT = Path(__file__).resolve().parents[2]
BOOK_DIR = REPO_ROOT / "03_Buch"
EXPORT_DIR = REPO_ROOT / "04_Export"
REFERENCE_DOCX = EXPORT_DIR / "reference.docx"
OUTPUT_DOCX = EXPORT_DIR / "Fremdfirmenkoordination.docx"
COMBINED_MD = EXPORT_DIR / "combined.md"

BOOK_TITLE = "Fremdfirmenkoordination"
BOOK_SUBTITLE = "Rechtliche Grundlagen und praktische Umsetzung\nin Bauwesen und Produktion"
BOOK_AUTHORS = ["Stapcon GmbH"]
BOOK_DATE = "2025"

# Kapitelreihenfolge mit Teilüberschriften
CHAPTERS = [
    {
        "part": None,
        "files": [],
        "title_page": True,
    },
    {
        "part": "Teil I – Grundlagen",
        "files": [
            "01_Grundlagen/01_Einleitung.md",
            "01_Grundlagen/02_Begriffe_und_Definitionen.md",
            "01_Grundlagen/03_Bedeutung_und_Risiken.md",
        ],
    },
    {
        "part": "Teil II – Rechtliche Grundlagen",
        "files": [
            "02_Rechtliche_Grundlagen/01_EU-Recht.md",
            "02_Rechtliche_Grundlagen/02_Deutsches_Arbeitsschutzrecht.md",
            "02_Rechtliche_Grundlagen/03_Zivilrecht.md",
            "02_Rechtliche_Grundlagen/04_Strafrecht.md",
            "02_Rechtliche_Grundlagen/05_Ordnungswidrigkeiten_und_Aufsicht.md",
        ],
    },
    {
        "part": "Teil III – Rollen und Verantwortlichkeiten",
        "files": [
            "03_Rollen_und_Verantwortlichkeiten/01_Auftraggeber_Betreiber.md",
            "03_Rollen_und_Verantwortlichkeiten/02_Auftragnehmer_Fremdfirma.md",
            "03_Rollen_und_Verantwortlichkeiten/03_Koordinator.md",
            "03_Rollen_und_Verantwortlichkeiten/04_SiFa_und_Betriebsarzt.md",
        ],
    },
    {
        "part": "Teil IV – Operative Prozesse",
        "files": [
            "04_Operative_Prozesse/01_Vergabe_und_Qualifizierung.md",
            "04_Operative_Prozesse/02_Gefährdungsbeurteilung.md",
            "04_Operative_Prozesse/03_Sicherheitsunterweisung.md",
            "04_Operative_Prozesse/04_Erlaubnisscheinverfahren.md",
            "04_Operative_Prozesse/05_Überwachung_und_Kontrolle.md",
            "04_Operative_Prozesse/06_Unfall_und_Schadensmanagement.md",
        ],
    },
    {
        "part": "Teil V – Branchenspezifische Anforderungen",
        "files": [
            "05_Branchenspezifisch/01_Bauwesen_und_Baustellenkoordination.md",
            "05_Branchenspezifisch/02_Produktion_und_Fertigung.md",
        ],
    },
    {
        "part": "Teil VI – Praktische Werkzeuge",
        "files": [
            "06_Praktische_Werkzeuge/01_Checklisten.md",
            "06_Praktische_Werkzeuge/02_Musterformulare.md",
            "06_Praktische_Werkzeuge/03_Mustervertragsklauseln.md",
            "06_Praktische_Werkzeuge/04_Digitalisierung.md",
        ],
    },
    {
        "part": "Anhang",
        "files": [
            "07_Anhang/01_Gesetzestexte_Referenztabelle.md",
            "07_Anhang/02_Literaturverzeichnis.md",
            "07_Anhang/03_Stichwortverzeichnis.md",
        ],
    },
]


# ---------------------------------------------------------------------------
# Schritt 1: Reference-DOCX mit professionellen Stilen erstellen
# ---------------------------------------------------------------------------

DARK_BLUE = RGBColor(0x00, 0x33, 0x66)
MID_BLUE  = RGBColor(0x00, 0x56, 0x96)
LIGHT_BLUE = RGBColor(0x1F, 0x78, 0xB4)
DARK_GRAY = RGBColor(0x33, 0x33, 0x33)


def set_font(run_or_paragraph, name="Calibri", size_pt=11, bold=False,
             italic=False, color=None):
    rpr = run_or_paragraph.font if hasattr(run_or_paragraph, "font") else None
    if rpr is None:
        return
    rpr.name = name
    rpr.size = Pt(size_pt)
    rpr.bold = bold
    rpr.italic = italic
    if color:
        rpr.color.rgb = color


def set_para_format(para, left_cm=0, space_before_pt=6, space_after_pt=6,
                    line_spacing_pt=None, keep_with_next=False,
                    page_break_before=False, alignment=None):
    pf = para.paragraph_format
    pf.left_indent = Cm(left_cm)
    pf.space_before = Pt(space_before_pt)
    pf.space_after = Pt(space_after_pt)
    if line_spacing_pt:
        from docx.shared import Pt as _Pt
        pf.line_spacing = _Pt(line_spacing_pt)
    pf.keep_with_next = keep_with_next
    pf.page_break_before = page_break_before
    if alignment is not None:
        pf.alignment = alignment


def add_section_break(doc, break_type="nextPage"):
    """Fügt einen Abschnittswechsel ein (für Kopf-/Fußzeile je Teil)."""
    new_section = doc.add_section()
    return new_section


def _set_heading_style(style, size_pt, color, bold=True, space_before=12,
                        space_after=6, keep_with_next=True,
                        page_break_before=False, italic=False):
    font = style.font
    font.name = "Calibri"
    font.size = Pt(size_pt)
    font.bold = bold
    font.italic = italic
    font.color.rgb = color
    pf = style.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.keep_with_next = keep_with_next
    pf.page_break_before = page_break_before


def create_reference_docx(path: Path):
    doc = Document()

    # Seitenränder: A4
    for section in doc.sections:
        section.page_width  = Cm(21)
        section.page_height = Cm(29.7)
        section.left_margin   = Cm(3.0)
        section.right_margin  = Cm(2.5)
        section.top_margin    = Cm(2.5)
        section.bottom_margin = Cm(2.0)
        section.header_distance = Cm(1.25)
        section.footer_distance = Cm(1.25)

    styles = doc.styles

    # Normal (Fließtext)
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = DARK_GRAY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after  = Pt(8)
    normal.paragraph_format.line_spacing = Pt(15)

    # Überschriften
    _set_heading_style(styles["Heading 1"], 18, DARK_BLUE,
                       space_before=24, space_after=12, page_break_before=True)
    _set_heading_style(styles["Heading 2"], 14, MID_BLUE,
                       space_before=18, space_after=8)
    _set_heading_style(styles["Heading 3"], 12, LIGHT_BLUE,
                       space_before=12, space_after=6)
    _set_heading_style(styles["Heading 4"], 11, DARK_GRAY,
                       bold=False, italic=True,
                       space_before=8, space_after=4)

    # Titel-Stil
    title_style = styles["Title"]
    title_style.font.name = "Calibri"
    title_style.font.size = Pt(28)
    title_style.font.bold = True
    title_style.font.color.rgb = DARK_BLUE
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_style.paragraph_format.space_before = Pt(72)
    title_style.paragraph_format.space_after  = Pt(12)

    # Untertitel
    subtitle_style = styles["Subtitle"]
    subtitle_style.font.name = "Calibri"
    subtitle_style.font.size = Pt(14)
    subtitle_style.font.color.rgb = MID_BLUE
    subtitle_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_style.paragraph_format.space_before = Pt(0)
    subtitle_style.paragraph_format.space_after  = Pt(48)

    # Aufzählungen
    for list_style_name in ("List Bullet", "List Bullet 2", "List Bullet 3"):
        if list_style_name in styles:
            s = styles[list_style_name]
            s.font.name = "Calibri"
            s.font.size = Pt(11)
            s.paragraph_format.space_before = Pt(2)
            s.paragraph_format.space_after  = Pt(2)
            s.paragraph_format.left_indent  = Cm(1.0)

    for list_style_name in ("List Number", "List Number 2"):
        if list_style_name in styles:
            s = styles[list_style_name]
            s.font.name = "Calibri"
            s.font.size = Pt(11)
            s.paragraph_format.space_before = Pt(2)
            s.paragraph_format.space_after  = Pt(2)
            s.paragraph_format.left_indent  = Cm(1.0)

    # Blockzitate (für Querverweise)
    if "Intense Quote" in styles:
        iq = styles["Intense Quote"]
        iq.font.name = "Calibri"
        iq.font.size = Pt(10)
        iq.font.italic = True
        iq.font.color.rgb = MID_BLUE
        iq.paragraph_format.left_indent = Cm(1.0)
        iq.paragraph_format.space_before = Pt(4)
        iq.paragraph_format.space_after  = Pt(4)

    # Kopf- und Fußzeile (erste Sektion)
    section = doc.sections[0]
    header = section.header
    hp = header.paragraphs[0]
    hp.text = BOOK_TITLE
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.style.font.name = "Calibri"
    hp.style.font.size = Pt(9)
    hp.style.font.color.rgb = MID_BLUE

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    # Automatische Seitenzahl via XML
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    fp.style.font.name = "Calibri"
    fp.style.font.size = Pt(9)
    fp.style.font.color.rgb = DARK_GRAY

    doc.save(str(path))
    print(f"  → Reference-DOCX erstellt: {path.name}")


# ---------------------------------------------------------------------------
# Schritt 2: Markdown bereinigen und zusammenführen
# ---------------------------------------------------------------------------

# Muster für Navigationszeilen (> *Querverweise:*)
NAV_LINE_RE = re.compile(r"^>\s*\*Querverweise.*$", re.MULTILINE)
# Kapitelbeschriftungen (**Teil I: ...**  /  **Kapitel 1: ...**)
KAPITEL_LABEL_RE = re.compile(r"^\*\*Teil\s.*$|^\*\*Kapitel\s.*$", re.MULTILINE)
# Einzelne Leerzeilen nach Querverweisblock bereinigen
TRIPLE_BLANK_RE = re.compile(r"\n{3,}", re.MULTILINE)


def clean_markdown(text: str) -> str:
    """Entfernt interne Navigationselemente aus dem Markdown."""
    # Navigations-Blockquotes entfernen
    text = NAV_LINE_RE.sub("", text)
    # Kapitelbeschriftungen entfernen (werden aus Dateinamen/Struktur neu gesetzt)
    text = KAPITEL_LABEL_RE.sub("", text)
    # Mehrfache Leerzeilen auf max. 2 reduzieren
    text = TRIPLE_BLANK_RE.sub("\n\n", text)
    return text.strip()


def extract_chapter_title(text: str) -> tuple[str, str]:
    """
    Extrahiert den Kapiteltitel aus dem ersten `#`-Heading.
    Gibt (title, rest_of_text) zurück.
    """
    lines = text.split("\n")
    title = ""
    rest_start = 0
    for i, line in enumerate(lines):
        stripped = line.strip()
        if stripped.startswith("# "):
            title = stripped[2:].strip()
            rest_start = i + 1
            break
    rest = "\n".join(lines[rest_start:]).strip()
    return title, rest


def build_combined_markdown() -> str:
    """Kombiniert alle Kapitel zu einem Markdown-Dokument."""
    parts = []

    # YAML-Front Matter
    authors_yaml = "\n".join(f'  - "{a}"' for a in BOOK_AUTHORS)
    parts.append(f"""---
title: "{BOOK_TITLE}"
subtitle: "{BOOK_SUBTITLE.replace(chr(10), ' ')}"
author:
{authors_yaml}
date: "{BOOK_DATE}"
lang: de-DE
toc: true
toc-depth: 3
numbersections: true
---
""")

    for chapter_group in CHAPTERS:
        if chapter_group.get("title_page"):
            continue

        part_title = chapter_group.get("part", "")
        if part_title:
            # Teilüberschrift als Heading 1 (ohne Nummerierung durch \unnumbered)
            parts.append(f"\n\n# {part_title} {{.unnumbered}}\n\n")

        for rel_path in chapter_group.get("files", []):
            md_path = BOOK_DIR / rel_path
            if not md_path.exists():
                print(f"  WARNUNG: Datei nicht gefunden: {rel_path}")
                continue

            raw = md_path.read_text(encoding="utf-8")
            cleaned = clean_markdown(raw)
            chapter_title, body = extract_chapter_title(cleaned)

            # Kapitelüberschrift als Heading 2
            if chapter_title:
                # Langen Buchtitel kürzen, wenn er mit dem Haupttitel übereinstimmt
                if chapter_title.startswith("Fremdfirmenkoordination"):
                    # Überschrift aus dem Dateinamen ableiten
                    chapter_title = md_path.stem.replace("_", " ")
                    chapter_title = re.sub(r"^\d+\s*", "", chapter_title).strip()
                parts.append(f"\n\n## {chapter_title}\n\n")

            parts.append(body)
            parts.append("\n\n")

    return "\n".join(parts)


# ---------------------------------------------------------------------------
# Hauptprogramm
# ---------------------------------------------------------------------------

def main():
    print("=" * 60)
    print("  Buchgenerierung: Fremdfirmenkoordination.docx")
    print("=" * 60)

    # Ausgabeverzeichnis anlegen
    EXPORT_DIR.mkdir(exist_ok=True)

    # 1. Reference-DOCX erstellen
    print("\n[1/3] Erstelle Formatvorlage (reference.docx) …")
    create_reference_docx(REFERENCE_DOCX)

    # 2. Markdown zusammenführen
    print("\n[2/3] Verarbeite Kapitel …")
    combined = build_combined_markdown()
    COMBINED_MD.write_text(combined, encoding="utf-8")
    line_count = combined.count("\n")
    print(f"  → {line_count} Zeilen kombiniertes Markdown")

    # 3. Pandoc-Konvertierung
    print("\n[3/3] Pandoc-Konvertierung → DOCX …")
    extra_args = [
        "--toc",
        "--toc-depth=3",
        f"--reference-doc={REFERENCE_DOCX}",
        "--from=markdown+smart+pipe_tables+fenced_code_blocks",
        "--standalone",
        "-V", "geometry:a4paper",
        "-V", "lang=de-DE",
    ]
    pypandoc.convert_file(
        str(COMBINED_MD),
        "docx",
        outputfile=str(OUTPUT_DOCX),
        extra_args=extra_args,
    )

    size_kb = OUTPUT_DOCX.stat().st_size // 1024
    print(f"\n✓ Fertig: {OUTPUT_DOCX}")
    print(f"  Dateigröße: {size_kb} KB")
    print(f"\n  Öffnen mit: libreoffice \"{OUTPUT_DOCX}\"")


if __name__ == "__main__":
    main()
